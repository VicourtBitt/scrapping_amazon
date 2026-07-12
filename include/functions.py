from selenium import webdriver
from selenium.webdriver.remote.webdriver import WebDriver
from webdriver_manager.chrome import ChromeDriverManager

from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from include.performance import perf_tracker

from datetime import datetime
import logging
import time
import re

_logger = logging.getLogger(__name__)

@perf_tracker
def setup_chrome_webdriver(headless: bool = True):
    """
    Automatically setups the Chrome Webdriver if it's not found in the OS.

    :param headless: It the instance should run in headless mode
    :return: driver
    """
    try:
        chrome_options = Options()
        if headless:
            chrome_options.add_argument("--headless")
            _logger.info("Chrome now will run in headless mode")

        # Install Chrome Drive if the system does not have it
        driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=chrome_options)
        _logger.info("Chrome WebDriver has now been initialized.")
        return driver

    except Exception as e:
        _logger.error(f"An error has been found while trying to setup the Chrome WebDriver. \n {e}", exc_info=True)

def safe_extract(element, by, selector, use_text_content=True, attribute=None):
    try:
        target = element.find_element(by, selector)
        if attribute:
            return target.get_attribute(attribute).strip()
        if use_text_content:
            return target.get_attribute("textContent").strip()
        return target.text.strip()
    except Exception:
        return "N/A"

@perf_tracker
def extract_best_sellers(driver: WebDriver, url: str, max_items: int = 6):
    best_sellers = []
    driver.get(url)

    wait = WebDriverWait(driver, 10)

    hamburger_btn = wait.until(
        EC.element_to_be_clickable((By.ID, "nav-hamburger-menu"))
    )
    hamburger_btn.click()

    # 2. Click "Ver tudo" (See all departments)
    # Using text and class combination to avoid hidden duplicate elements in the DOM
    see_all_btn = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//a[contains(@class, 'hmenu-compressed-btn') and contains(., 'Ver tudo')]",
            )
        )
    )
    # Sometimes Amazon menus require a slight scroll within the menu container
    driver.execute_script("arguments[0].scrollIntoView(true);", see_all_btn)
    see_all_btn.click()

    # 3. Click "Computadores e Informática"
    # Using the exact text inside the hmenu-item
    comp_menu_btn = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//a[contains(@class, 'hmenu-item') and contains(., 'Computadores e Informática')]",
            )
        )
    )
    driver.execute_script("arguments[0].scrollIntoView(true);", comp_menu_btn)
    comp_menu_btn.click()

    # 4. Click the first <li><a> link inside the "Computadores e Informática" section
    # Targets the first list item ('li[1]/a') inside the specific category section
    first_category_link = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//section[@aria-labelledby='Computadores e Informática']//ul/li[1]/a",
            )
        )
    )
    driver.execute_script("arguments[0].click();", first_category_link)

    carousel_xpath = f"//div[@data-reftag='brpc-best-sellers-dossier_0']//ol[contains(@class, 'a-carousel')]/li[position() <= {max_items}]"

    carousel_items = wait.until(
        EC.presence_of_all_elements_located((By.XPATH, carousel_xpath))
    )

    for index, item in enumerate(carousel_items, start=1):
        if len(best_sellers) >= max_items:
            break

        driver.execute_script(
            "arguments[0].scrollIntoView({behavior: 'instant', block: 'nearest', inline: 'center'});",
            item,
        )

        # Give Amazon's lazy-loader 0.5 seconds to populate images/prices if it was offscreen
        time.sleep(0.5)

        # NAME
        name = safe_extract(
            item, By.XPATH, ".//span[contains(@class, 'dcl-truncate dcl-product-title')]"
        )
        if name == "N/A" or not name:
            try:
                name = (
                    item.find_element(By.XPATH, ".//img")
                    .get_attribute("alt")
                    .strip()
                )
            except Exception:
                name = "Title not found"

        # PRICE
        price = safe_extract(
            item,
            By.XPATH,
            ".//span[contains(@class, 'a-price')]/span[contains(@class, 'a-offscreen')]",
            use_text_content=True,
        )

        # RATING
        rating = safe_extract(
            item,
            By.XPATH,
            ".//span[contains(@class, 'a-icon-alt')]",
            use_text_content=True,
        )

        # RATING COUNT (Now uses textContent automatically!)
        rating_count = safe_extract(
            item,
            By.XPATH,
            ".//span[contains(@class, 'dcl-product-rating-count')]",
        )

        link = safe_extract(item, By.XPATH, ".//a[@href]", attribute="href")

        best_sellers.append(
            {
                "position": index,
                "name": name,
                "price": price,
                "rating": rating,
                "rating_count": rating_count,
                "link": link
            }
        )
        print(f"[{index}] {name[:40]}... | {price} | {rating} ({rating_count})")
    return best_sellers

def clean_price(price_str):
    """Converts Brazilian price string 'R$ 1.056,00' to float 1056.00 for Excel math."""
    if not price_str or price_str == "N/A":
        return None
    try:
        cleaned = (
            price_str.replace("R$", "")
            .replace(".", "")
            .replace(" ", "")
            .replace(",", ".")
            .strip()
        )
        return float(cleaned)
    except Exception:
        return price_str


def clean_rating(rating_str):
    """Converts '4,8 de 5 estrelas' to float 4.8."""
    if not rating_str or rating_str == "N/A":
        return None
    try:
        match = re.search(r"(\d+[,.]\d+|\d+)", rating_str)
        if match:
            return float(match.group(1).replace(",", "."))
        return rating_str
    except Exception:
        return rating_str


def clean_count(count_str):
    """Converts '(1.420)' or '1.420' to integer 1420."""
    if not count_str or count_str == "N/A":
        return None
    try:
        cleaned = re.sub(r"[^\d]", "", count_str)
        return int(cleaned) if cleaned else None
    except Exception:
        return count_str

@perf_tracker
def export_to_excel(results, filename="amazon_mais_vendidos.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mais Vendidos - Informática"
    ws.views.sheetView[0].showGridLines = True

    # --- Color Palette & Styles ---
    primary_dark = "2C3E50"  # Dark slate blue for main title
    header_fill_hex = "34495e"  # Table header fill
    zebra_fill_hex = "F8F9FA"  # Alternating row fill
    border_color = "D3D3D3"  # Light grey borders

    title_font = Font(name="Segoe UI", size=16, bold=True, color=primary_dark)
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="565959")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F1111")
    bold_data_font = Font(name="Segoe UI", size=10, bold=True, color="0F1111")
    link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")

    header_fill = PatternFill(
        start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid"
    )
    zebra_fill = PatternFill(
        start_color=zebra_fill_hex, end_color=zebra_fill_hex, fill_type="solid"
    )
    white_fill = PatternFill(
        start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
    )
    summary_fill = PatternFill(
        start_color="EAEDED", end_color="EAEDED", fill_type="solid"
    )

    thin_side = Side(border_style="thin", color=border_color)
    double_side = Side(border_style="double", color="34495e")
    top_thick_side = Side(border_style="medium", color="34495e")

    cell_border = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    summary_border = Border(
        top=top_thick_side, bottom=double_side, left=thin_side, right=thin_side
    )

    # --- Row 1: Title ---
    num_items = len(results)
    ws["A1"] = f"Top {num_items} Mais Vendidos — Computadores e Informática"
    ws["A1"].font = title_font

    # --- Row 2: Extraction Date ---
    current_time = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")
    ws["A2"] = f"Data de Extração: {current_time} | Fonte: Amazon Brasil"
    ws["A2"].font = subtitle_font

    # --- Row 4: Table Headers ---
    headers = [
        "Posição",
        "Produto / Título do Item",
        "Preço (R$)",
        "Avaliação (0-5)",
        "Total de Avaliações",
        "Link do Produto",
    ]
    start_row = 4

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = cell_border
        cell.alignment = Alignment(
            horizontal="center" if col_idx != 2 else "left",
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[start_row].height = 24

    # --- Rows 5+: Data Insertion ---
    for i, item in enumerate(results):
        current_row = start_row + 1 + i
        ws.row_dimensions[current_row].height = 20
        row_fill = zebra_fill if i % 2 == 1 else white_fill

        # 1. Posição
        c_pos = ws.cell(row=current_row, column=1, value=item.get("position", i + 1))
        c_pos.alignment = Alignment(horizontal="center", vertical="center")

        # 2. Produto
        c_name = ws.cell(row=current_row, column=2, value=item.get("name", "N/A"))
        c_name.alignment = Alignment(horizontal="left", vertical="center")

        # 3. Preço
        val_price = clean_price(item.get("price"))
        c_price = ws.cell(
            row=current_row,
            column=3,
            value=val_price if val_price is not None else item.get("price", "N/A"),
        )
        c_price.alignment = Alignment(horizontal="right", vertical="center")
        if isinstance(val_price, float):
            c_price.number_format = '"R$" #,##0.00'

        # 4. Avaliação
        val_rating = clean_rating(item.get("rating"))
        c_rating = ws.cell(
            row=current_row,
            column=4,
            value=(
                val_rating
                if val_rating is not None
                else item.get("rating", "N/A")
            ),
        )
        c_rating.alignment = Alignment(horizontal="center", vertical="center")
        if isinstance(val_rating, float):
            c_rating.number_format = '0.0 "★"'

        # 5. Total de Avaliações
        val_count = clean_count(item.get("rating_count"))
        c_count = ws.cell(
            row=current_row,
            column=5,
            value=(
                val_count
                if val_count is not None
                else item.get("rating_count", "N/A")
            ),
        )

        c_count.alignment = Alignment(horizontal="right", vertical="center")
        if isinstance(val_count, int):
            c_count.number_format = "#,##0"

        # 6. Link do Produto
        val_link = item.get("link", "N/A")
        c_link = ws.cell(
            row=current_row,
            column=6,
            value="Ver na Amazon" if val_link != "N/A" else "N/A",
        )

        c_link.alignment = Alignment(horizontal="center", vertical="center")
        if val_link != "N/A":
            c_link.hyperlink = val_link
            c_link.font = link_font

        # Apply borders and fills to all cells in the row
        for col_idx in range(1, 7):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = (
                data_font if col_idx != 6 or val_link == "N/A" else link_font
            )
            cell.fill = row_fill
            cell.border = cell_border

    # --- Summary Row (Average & Totals) ---
    end_row = start_row + len(results)
    summary_row = end_row + 1
    ws.row_dimensions[summary_row].height = 22

    ws.cell(row=summary_row, column=1, value="Média / Total").alignment = (
        Alignment(horizontal="center", vertical="center")
    )
    ws.cell(
        row=summary_row,
        column=2,
        value="Resumo Estatístico dos Itens Extraídos",
    ).alignment = Alignment(horizontal="left", vertical="center")

    # Excel Formulas
    c_avg_price = ws.cell(
        row=summary_row, column=3, value=f"=AVERAGE(C{start_row+1}:C{end_row})"
    )
    c_avg_price.number_format = '"R$" #,##0.00'
    c_avg_price.alignment = Alignment(horizontal="right", vertical="center")

    c_avg_rating = ws.cell(
        row=summary_row, column=4, value=f"=AVERAGE(D{start_row+1}:D{end_row})"
    )
    c_avg_rating.number_format = '0.0 "★"'
    c_avg_rating.alignment = Alignment(horizontal="center", vertical="center")

    c_tot_reviews = ws.cell(
        row=summary_row, column=5, value=f"=SUM(E{start_row+1}:E{end_row})"
    )
    c_tot_reviews.number_format = "#,##0"
    c_tot_reviews.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, 7):
        cell = ws.cell(row=summary_row, column=col_idx)
        cell.font = bold_data_font
        cell.fill = summary_fill
        cell.border = summary_border

    # --- Layout Polish ---
    ws.freeze_panes = f"A{start_row+1}"
    ws.auto_filter.ref = f"A{start_row}:F{end_row}"

    # Auto-adjust column widths with padding
    column_min_widths = {1: 12, 2: 50, 3: 18, 4: 18, 5: 22, 6:18}
    for col_idx in range(1, 7):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(start_row, end_row + 2):
            val = str(ws.cell(row=row, column=col_idx).value or "")
            if val.startswith("="):
                val = "R$ 0000,00"  # Estimate display length for formulas
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(
            max_len + 4, column_min_widths[col_idx]
        )

    wb.save(filename)
    print(f"\nExcel file successfully generated and saved as: {filename}")
    return filename